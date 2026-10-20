import sys, time, json, os, re
from pathlib import Path
import pypdfium2 as pdfium
import docx, openpyxl, olefile
ROOTS=[r for r in os.environ.get('BENCH_ROOTS', '').split(';') if r]  # dossiers de test privés, ex. BENCH_ROOTS="D:\a;D:\b"
def pdf(p):
    d=pdfium.PdfDocument(p); n=len(d); t=[]
    for i in range(n):
        pg=d[i]; t.append(pg.get_textpage().get_text_range()); pg.close()
    d.close(); return "\n".join(t), n
def dx(p):
    d=docx.Document(p); t=[x.text for x in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows: t.append(" | ".join(c.text for c in r.cells))
    return "\n".join(t), 1
def xl(p):
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True); t=[]
    for ws in wb.worksheets:
        t.append("## "+ws.title)
        for r in ws.iter_rows(values_only=True):
            v=[str(c) for c in r if c is not None]
            if v: t.append(" | ".join(v))
    return "\n".join(t), len(wb.worksheets)
def doc(p):
    o=olefile.OleFileIO(p)
    if not o.exists('WordDocument'): return "",1
    b=o.openstream('WordDocument').read()
    s=b.decode('latin-1',errors='ignore'); s2=b.decode('utf-16le',errors='ignore')
    pick=max([s,s2],key=lambda x:len(re.findall(r'[A-Za-zÀ-ÿ]{3,}',x)))
    return " ".join(re.findall(r'[\w\sÀ-ÿ.,;:!?\'-]{20,}',pick)),1
H={'.pdf':pdf,'.docx':dx,'.xlsx':xl,'.doc':doc}
res=[]
for r in ROOTS:
    for p in Path(r).rglob('*'):
        e=p.suffix.lower()
        if e in H and p.is_file():
            t0=time.perf_counter()
            try: txt,n=H[e](str(p)); err=None
            except Exception as ex: txt,n,err="",0,str(ex)[:80]
            res.append(dict(path=str(p),ext=e,mb=p.stat().st_size/1e6,sec=time.perf_counter()-t0,chars=len(txt),units=n,err=err))
json.dump(res,open('extract_results.json','w'))
from collections import defaultdict
g=defaultdict(list)
for x in res: g[x['ext']].append(x)
for e,l in g.items():
    sec=sum(x['sec'] for x in l); ch=sum(x['chars'] for x in l)
    empty=[x for x in l if x['chars']<50 and not x['err']]; errs=[x for x in l if x['err']]
    print(f"{e}: n={len(l)} time={sec:.1f}s avg={sec/len(l)*1000:.0f}ms chars={ch/1e6:.2f}M empty(<50ch)={len(empty)} errors={len(errs)}")
    if e=='.pdf': print("  pages total:",sum(x['units'] for x in l))
    for x in errs[:3]: print("  ERR",Path(x['path']).name,x['err'])
